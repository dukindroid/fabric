# -*- coding: utf-8 -*-
"""
Genera las dos versiones de CV en formato DOCX a partir de los archivos .md.
Estilo inspirado en curr-2026.pdf: nombre grande y centrado, datos de
contacto, secciones en MAYÚSCULAS, viñetas y negritas.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DARK = RGBColor(0x1F, 0x38, 0x64)   # azul oscuro para secciones

BASE_DIR = Path(__file__).parent

SOURCES = {
    "CV-Holcim-Operario-Maquinaria-Pesada.md": "CV-Holcim-Operario-Maquinaria-Pesada.docx",
    "CV-Bodesa-Especialista-Integridad-Producto.md": "CV-Bodesa-Especialista-Integridad-Producto.docx",
}

CONTACT_MARKERS = ("dukintosh@gmail.com", "LinkedIn")


def add_run_with_bold(paragraph, text):
    """Añade texto a un párrafo respetando **negritas** inline."""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def build(md_path, docx_path):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)

    name_done = False
    contact_done = False

    for raw in lines:
        line = raw.strip()
        if not line or line.startswith("---"):
            continue

        # Nombre principal
        if line.startswith("# ") and not name_done:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(" ".join(line[2:].split()))  # normaliza espacios dobles
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = DARK
            p.paragraph_format.space_after = Pt(2)
            name_done = True
            continue

        # Línea de contacto (contiene email/LinkedIn)
        if not contact_done and any(m in line for m in CONTACT_MARKERS):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_bold(p, line)
            p.paragraph_format.space_after = Pt(4)
            contact_done = True
            continue

        # Sección principal (## MAYÚSCULAS)
        if line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = DARK
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            continue

        # Puesto / subsección en experiencia (###)
        if line.startswith("### "):
            p = doc.add_paragraph()
            add_run_with_bold(p, line[4:].strip())
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10.5)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            continue

        # Viñetas
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_run_with_bold(p, line[2:].strip())
            p.paragraph_format.space_after = Pt(1)
            continue

        # Párrafo normal (perfil, fechas en negritas, subencabezados de habilidades)
        p = doc.add_paragraph()
        add_run_with_bold(p, line)
        p.paragraph_format.space_after = Pt(2)

    doc.save(docx_path)
    print(f"OK -> {docx_path.name}")


def main():
    for md, docx in SOURCES.items():
        build(BASE_DIR / md, BASE_DIR / docx)


if __name__ == "__main__":
    main()